"""
Export a tailored resume (JSON) to a Word document (.docx) or PDF.

This is a MANUAL, opt-in step — deliberately separate from main.py's
pipeline. Review the tailored JSON and guardrail/eval report first; only
export to a document once you're actually happy with the content. Nothing
in the tailoring pipeline calls this automatically.

Usage:
    python export_resume.py --input output/tailoring_report.json --format both
    python export_resume.py --input sample_data/base_resume.json --format docx --out output/my_resume

--input accepts either:
  - a full TailoringReport JSON (as saved by main.py) — the tailored_resume
    field is pulled out automatically
  - a raw Resume JSON (e.g. sample_data/base_resume.json), useful for
    exporting your untailored base resume too

Layout notes: this builds a tight, single-column resume layout (narrow
margins, minimal paragraph spacing, bold-caps section headers with a thin
divider line underneath) rather than relying on Word's default document
template — Word's defaults alone (1" margins, ~8pt paragraph spacing,
colored Heading styles) waste enough space that a 1-page resume can spill
onto a second page even with the right amount of content.
"""

import argparse
import json
from pathlib import Path

from models import Resume
from style_extractor import extract_style_profile, map_font_to_reportlab, DEFAULT_PROFILE


def load_resume_from_file(path: str) -> Resume:
    with open(path) as f:
        data = json.load(f)
    if "tailored_resume" in data:  # it's a TailoringReport
        return Resume.model_validate(data["tailored_resume"])
    return Resume.model_validate(data)  # it's a raw Resume


def _contact_line(resume: Resume) -> str:
    bits = [b for b in [resume.email, resume.phone, resume.location] if b]
    links = [resume.links.linkedin, resume.links.github, resume.links.portfolio] + resume.links.other
    bits += [l for l in links if l]
    return " | ".join(bits)


# ---------------------------------------------------------------------------
# DOCX export
# ---------------------------------------------------------------------------

def export_docx(resume: Resume, out_path: str, style_profile: dict = None):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    style = style_profile or DEFAULT_PROFILE
    doc = Document()

    # --- Tight page margins (Word's 1" default eats a lot of a 1-page budget) ---
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    def _tighten(paragraph, space_before=0, space_after=2, line_spacing=1.0):
        pf = paragraph.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = line_spacing

    def _add_bottom_border(paragraph, size=6, color="000000"):
        """Adds a thin horizontal rule below a paragraph — used as a section divider."""
        p = paragraph._p
        pPr = p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), color)
        pBdr.append(bottom)
        pPr.append(pBdr)

    def add_section_heading(text):
        """Bold, all-caps section header with a divider line underneath —
        NOT Word's built-in Heading style, which is themed/colored by default
        and doesn't match typical resume conventions."""
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.name = style["heading_font"]
        run.font.size = Pt(style["heading_size"])
        run.font.color.rgb = RGBColor(0, 0, 0)
        _tighten(p, space_before=8, space_after=2)
        _add_bottom_border(p)
        return p

    # Apply body font/size to the whole doc's default style, so every plain
    # paragraph inherits it without extra per-run work
    normal = doc.styles["Normal"]
    normal.font.name = style["body_font"]
    normal.font.size = Pt(style["body_size"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    # "List Bullet" is a separate built-in style with its own spacing —
    # has to be tightened independently or bullets end up double-spaced
    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.name = style["body_font"]
    bullet_style.font.size = Pt(style["body_size"])
    bullet_style.paragraph_format.space_before = Pt(0)
    bullet_style.paragraph_format.space_after = Pt(1)
    bullet_style.paragraph_format.line_spacing = 1.0

    name_para = doc.add_paragraph()
    name_run = name_para.add_run(resume.name)
    name_run.bold = True
    name_run.font.name = style["name_font"]
    name_run.font.size = Pt(style["name_size"])
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tighten(name_para, space_after=2)

    contact_line = _contact_line(resume)
    if contact_line:
        contact_para = doc.add_paragraph(contact_line)
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tighten(contact_para, space_after=6)

    if resume.summary:
        add_section_heading("Summary")
        p = doc.add_paragraph(resume.summary)
        _tighten(p)

    if resume.skills:
        add_section_heading("Skills")
        for cat in resume.skills:
            cat_para = doc.add_paragraph()
            cat_para.add_run(f"{cat.category}: ").bold = True
            cat_para.add_run(", ".join(cat.skills))
            _tighten(cat_para)

    if resume.experience:
        add_section_heading("Experience")
        for exp in resume.experience:
            role_para = doc.add_paragraph()
            role_para.add_run(f"{exp.title} — {exp.company}").bold = True
            _tighten(role_para, space_before=4, space_after=0)
            date_loc = f"{exp.start_date} – {exp.end_date}"
            if exp.location:
                date_loc += f" | {exp.location}"
            meta_para = doc.add_paragraph()
            meta_run = meta_para.add_run(date_loc)
            meta_run.italic = True
            meta_run.font.size = Pt(max(style["body_size"] - 1, 8))
            _tighten(meta_para, space_after=2)
            for bullet in exp.bullets:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                _tighten(bp, space_after=1)

    if resume.projects:
        add_section_heading("Projects")
        for proj in resume.projects:
            proj_para = doc.add_paragraph()
            proj_para.add_run(proj.name).bold = True
            _tighten(proj_para, space_before=4)
            if proj.description:
                dp = doc.add_paragraph(proj.description)
                _tighten(dp)
            for bullet in proj.bullets:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                _tighten(bp, space_after=1)

    if resume.education:
        add_section_heading("Education")
        for edu in resume.education:
            edu_para = doc.add_paragraph()
            edu_para.add_run(edu.institution).bold = True
            _tighten(edu_para, space_before=4, space_after=0)
            date_range = " – ".join([d for d in [edu.start_date, edu.end_date] if d])
            dp = doc.add_paragraph(edu.degree + (f" | {date_range}" if date_range else ""))
            _tighten(dp)

    if resume.certifications:
        add_section_heading("Certifications")
        for cert in resume.certifications:
            line = cert.name
            extra = " — ".join([e for e in [cert.issuer, cert.date] if e])
            if extra:
                line += f" ({extra})"
            cp = doc.add_paragraph(line, style="List Bullet")
            _tighten(cp, space_after=1)

    doc.save(out_path)
    print(f"Saved DOCX to {out_path}")


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------

def export_pdf(resume: Resume, out_path: str, style_profile: dict = None):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.colors import HexColor, black
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable

    style = style_profile or DEFAULT_PROFILE
    body_font = map_font_to_reportlab(style["body_font"])
    heading_font = map_font_to_reportlab(style["heading_font"])
    name_font = map_font_to_reportlab(style["name_font"])
    body_font_bold = "Times-Bold" if body_font == "Times-Roman" else f"{body_font}-Bold"

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("NameStyle", parent=styles["Title"], alignment=TA_CENTER,
                                 fontName=name_font, fontSize=style["name_size"],
                                 spaceAfter=2, textColor=black)
    contact_style = ParagraphStyle("ContactStyle", parent=styles["Normal"], alignment=TA_CENTER,
                                    fontName=body_font, fontSize=9, spaceAfter=6)
    heading_style = ParagraphStyle("SectionHeading", parent=styles["Normal"], spaceBefore=8, spaceAfter=1,
                                    fontName=body_font_bold, fontSize=style["heading_size"], textColor=black)
    role_style = ParagraphStyle("RoleStyle", parent=styles["Normal"], fontName=body_font_bold, spaceBefore=4,
                                 spaceAfter=0, fontSize=style["body_size"])
    meta_style = ParagraphStyle("MetaStyle", parent=styles["Normal"], fontSize=max(style["body_size"] - 1, 8),
                                 textColor=HexColor("#555555"), fontName=body_font, spaceAfter=2)
    body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontName=body_font,
                                 fontSize=style["body_size"], spaceAfter=2, leading=style["body_size"] + 2)
    bullet_style = ParagraphStyle("BulletStyle", parent=body_style, spaceAfter=1)

    doc = SimpleDocTemplate(
        out_path, pagesize=letter,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch,
    )
    story = [Paragraph(resume.name, name_style)]

    contact_line = _contact_line(resume)
    if contact_line:
        story.append(Paragraph(contact_line, contact_style))

    def add_section_heading(text):
        story.append(Paragraph(text.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=black, spaceBefore=0, spaceAfter=3))

    if resume.summary:
        add_section_heading("Summary")
        story.append(Paragraph(resume.summary, body_style))

    if resume.skills:
        add_section_heading("Skills")
        for cat in resume.skills:
            story.append(Paragraph(f"<b>{cat.category}:</b> {', '.join(cat.skills)}", body_style))

    if resume.experience:
        add_section_heading("Experience")
        for exp in resume.experience:
            story.append(Paragraph(f"{exp.title} — {exp.company}", role_style))
            date_loc = f"{exp.start_date} – {exp.end_date}" + (f" | {exp.location}" if exp.location else "")
            story.append(Paragraph(date_loc, meta_style))
            if exp.bullets:
                items = [ListItem(Paragraph(b, bullet_style), spaceAfter=1) for b in exp.bullets]
                story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=14))

    if resume.projects:
        add_section_heading("Projects")
        for proj in resume.projects:
            story.append(Paragraph(proj.name, role_style))
            if proj.description:
                story.append(Paragraph(proj.description, body_style))
            if proj.bullets:
                items = [ListItem(Paragraph(b, bullet_style), spaceAfter=1) for b in proj.bullets]
                story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=14))

    if resume.education:
        add_section_heading("Education")
        for edu in resume.education:
            story.append(Paragraph(edu.institution, role_style))
            date_range = " – ".join([d for d in [edu.start_date, edu.end_date] if d])
            story.append(Paragraph(edu.degree + (f" | {date_range}" if date_range else ""), body_style))

    if resume.certifications:
        add_section_heading("Certifications")
        cert_lines = []
        for c in resume.certifications:
            extra = " — ".join([e for e in [c.issuer, c.date] if e])
            cert_lines.append(f"{c.name} ({extra})" if extra else c.name)
        items = [ListItem(Paragraph(c, bullet_style), spaceAfter=1) for c in cert_lines]
        story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=14))

    doc.build(story)
    print(f"Saved PDF to {out_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export a tailored resume to Word and/or PDF (manual step).")
    parser.add_argument("--input", required=True, help="Path to a TailoringReport or raw Resume JSON file")
    parser.add_argument("--format", choices=["docx", "pdf", "both"], required=True)
    parser.add_argument("--out", default="output/resume", help="Output path WITHOUT extension, e.g. output/resume")
    DEFAULT_TEMPLATE = "sample_data/master_resume_style.pdf"
    parser.add_argument(
        "--template",
        default=DEFAULT_TEMPLATE if Path(DEFAULT_TEMPLATE).exists() else None,
        help=f"Path to a sample resume (.docx or .pdf) whose fonts/sizes should be matched. "
             f"Defaults to '{DEFAULT_TEMPLATE}' (your master resume) if present — pass "
             f"--template none to skip style matching entirely and use the plain default look. "
             f"Only lifts font choices and sizes, not layout/columns/colors — see style_extractor.py.",
    )
    args = parser.parse_args()
    if args.template and args.template.lower() == "none":
        args.template = None

    resume = load_resume_from_file(args.input)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)

    style_profile = None
    if args.template:
        style_profile = extract_style_profile(args.template)
        print(f"Extracted style from {args.template}: {style_profile}")

    if args.format in ("docx", "both"):
        export_docx(resume, f"{args.out}.docx", style_profile=style_profile)
    if args.format in ("pdf", "both"):
        export_pdf(resume, f"{args.out}.pdf", style_profile=style_profile)
